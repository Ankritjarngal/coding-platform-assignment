#!/usr/bin/env python3
"""
Seminar PPT — Dark Modern / Keynote-Minimal theme.

Design research applied:
- SlideSpeak Monolith: near-black stage, huge type, monospace labels, hairlines, empty space
- Keynote Minimal: one idea per slide, oversized headlines, single accent
- Dark Modern (tech talks): #0E1116 / #E8EAED / #9AA4B2 / #4DA3FF
- Duarte slide:ology: contrast, hierarchy, whitespace, 3-second rule
- Apple keynote practice: sparse text, large type, no bullet walls

9 slides. Enough content to teach. Not corporate teal-bar template.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt, Emu


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Ankrit_Jarngal_2023BSE028_Seminar_PPT.pptx"
ASSETS = ROOT / "assets"

# Dark Modern palette
BG = RGBColor(0x0E, 0x11, 0x16)
PANEL = RGBColor(0x17, 0x1B, 0x22)
TEXT = RGBColor(0xE8, 0xEA, 0xED)
MUTED = RGBColor(0x9A, 0xA4, 0xB2)
FAINT = RGBColor(0x5B, 0x65, 0x72)
ACCENT = RGBColor(0x4D, 0xA3, 0xFF)
HAIR = RGBColor(0x2A, 0x31, 0x3B)

FONT = "Arial"          # widely available; Inter-like sans
FONT_MONO = "Consolas"  # monospace labels


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def textbox(slide, left, top, width, height, text, *, size=18, bold=False,
            color=TEXT, align=PP_ALIGN.LEFT, font=FONT, tracking=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return box


def hairline(slide, left, top, width):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = HAIR
    shape.line.fill.background()
    return shape


def accent_line(slide, left, top, width):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(2.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()
    return shape


def eyebrow(slide, text, left=Inches(0.7), top=Inches(0.35)):
    """Monospace uppercase whisper label — Monolith / Dark Modern."""
    textbox(
        slide, left, top, Inches(12), Inches(0.35),
        text.upper(),
        size=11, bold=True, color=ACCENT, font=FONT_MONO,
    )


def panel(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = HAIR
    shape.line.width = Pt(1)
    try:
        shape.adjustments[0] = 0.06
    except Exception:
        pass
    return shape


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, BG)
    return s


def chrome(slide, page, total=9):
    """Faint top brand row + slide number — Monolith structure."""
    textbox(slide, Inches(0.7), Inches(0.22), Inches(6), Inches(0.3),
            "AI-POWERED REVERSE ENGINEERING", size=10, color=FAINT, font=FONT_MONO)
    textbox(slide, Inches(11.2), Inches(0.22), Inches(1.4), Inches(0.3),
            f"{page:02d} / {total:02d}", size=10, color=FAINT, font=FONT_MONO, align=PP_ALIGN.RIGHT)
    hairline(slide, Inches(0.7), Inches(0.55), Inches(11.9))


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ------------------------------------------------------------------ 01 TITLE
    s = blank(prs)
    eyebrow(s, "seminar  ·  nit srinagar  ·  2025–26", top=Inches(1.6))
    textbox(s, Inches(0.7), Inches(2.2), Inches(12), Inches(1.6),
            "AI-Powered\nReverse Engineering",
            size=54, bold=True, color=TEXT, align=PP_ALIGN.LEFT)
    accent_line(s, Inches(0.7), Inches(4.15), Inches(1.6))
    textbox(s, Inches(0.7), Inches(4.5), Inches(11), Inches(0.5),
            "How AI accelerates binary analysis — with video games as the case study",
            size=18, color=MUTED)
    textbox(s, Inches(0.7), Inches(6.4), Inches(11), Inches(0.5),
            "Ankrit Jarngal   ·   2023BSE028   ·   CSE",
            size=14, color=FAINT, font=FONT_MONO)

    # ------------------------------------------------------------------ 02 WHAT IS RE
    s = blank(prs)
    chrome(s, 2)
    eyebrow(s, "01  background", top=Inches(0.85))
    textbox(s, Inches(0.7), Inches(1.4), Inches(11.8), Inches(1.4),
            "Reverse engineering recovers meaning\nfrom compiled software.",
            size=36, bold=True, color=TEXT)
    textbox(s, Inches(0.7), Inches(3.4), Inches(11.5), Inches(2.2),
            "You start with a finished binary — not source code.\n"
            "The goal is to understand structure and behavior,\n"
            "then modify it carefully when needed.",
            size=20, color=MUTED)
    textbox(s, Inches(0.7), Inches(6.2), Inches(11.5), Inches(0.5),
            "Not rewriting from scratch. Interpreting what already exists.",
            size=16, bold=True, color=ACCENT)

    # ------------------------------------------------------------------ 03 TRADITIONAL PIPELINE
    s = blank(prs)
    chrome(s, 3)
    eyebrow(s, "01  background", top=Inches(0.85))
    textbox(s, Inches(0.7), Inches(1.25), Inches(12), Inches(0.7),
            "The classical pipeline",
            size=32, bold=True, color=TEXT)
    textbox(s, Inches(0.7), Inches(1.95), Inches(12), Inches(0.4),
            "Disassemble → decompile → understand. Tools are strong. Humans are the bottleneck.",
            size=16, color=MUTED)
    s.shapes.add_picture(str(ASSETS / "diagram_traditional_re.png"),
                         Inches(0.5), Inches(2.6), width=Inches(12.3))
    textbox(s, Inches(0.7), Inches(6.35), Inches(12), Inches(0.5),
            "Manual naming and structure recovery is where most time is lost.",
            size=15, color=MUTED)

    # ------------------------------------------------------------------ 04 WHAT IS AI-ASSISTED
    s = blank(prs)
    chrome(s, 4)
    eyebrow(s, "02  ai-assisted re", top=Inches(0.85))
    textbox(s, Inches(0.7), Inches(1.35), Inches(12), Inches(1.2),
            "AI sits on top of classical tools.\nIt does not replace them.",
            size=34, bold=True, color=TEXT)

    # three sparse panels
    items = [
        ("LLM layer", "Rename, explain, and refine\npseudo-code into readable form."),
        ("Tool bridges", "Connect models to Ghidra / IDA\nso agents can query the binary."),
        ("Verify loop", "Propose → check against bytes\n→ retry until evidence holds."),
    ]
    for i, (h, d) in enumerate(items):
        left = Inches(0.7 + i * 4.1)
        panel(s, left, Inches(3.5), Inches(3.8), Inches(2.6))
        textbox(s, left + Inches(0.3), Inches(3.75), Inches(3.2), Inches(0.45),
                h, size=18, bold=True, color=ACCENT, font=FONT_MONO)
        textbox(s, left + Inches(0.3), Inches(4.4), Inches(3.2), Inches(1.4),
                d, size=15, color=MUTED)

    # ------------------------------------------------------------------ 05 HOW FASTER
    s = blank(prs)
    chrome(s, 5)
    eyebrow(s, "02  ai-assisted re", top=Inches(0.85))
    textbox(s, Inches(0.7), Inches(1.2), Inches(6.5), Inches(1.0),
            "What got faster",
            size=34, bold=True, color=TEXT)
    textbox(s, Inches(0.7), Inches(2.2), Inches(5.8), Inches(3.8),
            "Readable drafts in minutes,\nnot weeks of renaming.\n\n"
            "Agent loops automate\ntry → check → fix.\n\n"
            "Same tools. Compressed time.\nVerification still required.",
            size=18, color=MUTED)
    s.shapes.add_picture(str(ASSETS / "diagram_ai_assisted.png"),
                         Inches(7.0), Inches(1.6), width=Inches(5.7))

    # ------------------------------------------------------------------ 06 CASE STUDY
    s = blank(prs)
    chrome(s, 6)
    eyebrow(s, "03  case study", top=Inches(0.85))
    textbox(s, Inches(0.7), Inches(1.35), Inches(12), Inches(1.3),
            "Case study: modifying compiled\nvideo games without source.",
            size=34, bold=True, color=TEXT)
    textbox(s, Inches(0.7), Inches(3.2), Inches(11.5), Inches(2.8),
            "Games ship as closed, optimized binaries with thin official APIs.\n"
            "Deep changes need reverse engineering of the executable itself.\n\n"
            "Recent work combines systems from different games at runtime —\n"
            "not by rewriting engines, but by bridging two live compiled processes.\n"
            "AI made that glue work comparatively practical.",
            size=18, color=MUTED)

    # ------------------------------------------------------------------ 07 HOW IT WORKS
    s = blank(prs)
    chrome(s, 7)
    eyebrow(s, "03  case study", top=Inches(0.85))
    textbox(s, Inches(0.7), Inches(1.2), Inches(12), Inches(0.7),
            "How the bridge works",
            size=32, bold=True, color=TEXT)
    s.shapes.add_picture(str(ASSETS / "diagram_case_study.png"),
                         Inches(0.9), Inches(2.1), width=Inches(11.4))
    textbox(s, Inches(0.7), Inches(5.7), Inches(12), Inches(1.2),
            "Guest and host stay as real binaries. A thin bridge exchanges camera,\n"
            "collision, and events over shared memory — then composites one frame.\n"
            "AI drafts hooks and protocols. Classical RE validates the bytes.",
            size=16, color=MUTED)

    # ------------------------------------------------------------------ 08 LIMITS
    s = blank(prs)
    chrome(s, 8)
    eyebrow(s, "04  limits", top=Inches(0.85))
    textbox(s, Inches(0.7), Inches(1.25), Inches(12), Inches(0.7),
            "Faster is not unrestricted.",
            size=34, bold=True, color=TEXT)

    limits = [
        ("Correctness", "AI can invent plausible but wrong logic. Verify against the binary."),
        ("Hard defenses", "Obfuscation and packing still break many automated approaches."),
        ("Legal lines", "Copyright and license terms still apply to commercial assets."),
        ("Responsibility", "Serious work stays offline. Do not ship proprietary files."),
    ]
    for i, (h, d) in enumerate(limits):
        col = i % 2
        row = i // 2
        left = Inches(0.7 + col * 6.2)
        top = Inches(2.3 + row * 2.1)
        panel(s, left, top, Inches(5.9), Inches(1.85))
        textbox(s, left + Inches(0.35), top + Inches(0.35), Inches(5.2), Inches(0.4),
                h, size=16, bold=True, color=ACCENT, font=FONT_MONO)
        textbox(s, left + Inches(0.35), top + Inches(0.9), Inches(5.2), Inches(0.7),
                d, size=15, color=MUTED)

    # ------------------------------------------------------------------ 09 CLOSE
    s = blank(prs)
    chrome(s, 9)
    eyebrow(s, "05  takeaway", top=Inches(0.85))
    textbox(s, Inches(0.7), Inches(2.0), Inches(12), Inches(1.6),
            "AI compresses the reverse-engineering\nbottleneck. It does not erase it.",
            size=36, bold=True, color=TEXT)
    accent_line(s, Inches(0.7), Inches(4.0), Inches(1.6))
    textbox(s, Inches(0.7), Inches(4.4), Inches(11.5), Inches(1.0),
            "Classical tools measure the bytes. LLMs draft the meaning.\n"
            "Humans — and verification oracles — decide what ships.",
            size=18, color=MUTED)
    textbox(s, Inches(0.7), Inches(6.3), Inches(12), Inches(0.5),
            "Thank you   ·   Ankrit Jarngal  ·  2023BSE028  ·  CSE  ·  NIT Srinagar",
            size=14, color=FAINT, font=FONT_MONO)

    notes = [
        "Open with topic and identity. Keep it calm.",
        "Define RE for a CSE audience. Stress: not rewriting from scratch.",
        "Walk the classical pipeline. Point at Manual Analysis as the costly step.",
        "Define AI-assisted RE as a layer, not a replacement.",
        "Explain speed: renaming, agent loops, verification still required.",
        "Introduce games as closed binaries — avoid brand-name dumping.",
        "Explain host/guest/bridge simply using the diagram.",
        "Cover correctness, defenses, legal, responsibility.",
        "Close on the thesis. Invite questions.",
    ]
    for slide, note in zip(prs.slides, notes):
        slide.notes_slide.notes_text_frame.text = note

    assert len(prs.slides) == 9
    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides) — Dark Modern / Keynote-Minimal theme")


if __name__ == "__main__":
    build()
